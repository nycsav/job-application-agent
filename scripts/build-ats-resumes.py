#!/usr/bin/env python3
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "materials" / "resumes" / "optimized"

CONTACT = "New York, NY | 415-828-5282 | sav@ensopartners.co | linkedin.com/in/savbanerjee | ensolabs.ai | github.com/nycsav"

BASE_EXPERIENCE = [
    {
        "company": "Enso Labs",
        "title": "Founder & Principal, AI Transformation",
        "dates": "2020 - Present | New York, NY",
        "bullets": [
            "Built and operate a principal-led AI transformation studio serving enterprise clients across healthcare, finance, B2B technology, and advanced manufacturing.",
            "Own end-to-end engagement lifecycle with C-suite and senior client stakeholders: discovery, SOW, AI roadmap, operating model, architecture, build, deployment, governance, and value realization.",
            "Average 3-month time-to-first-value and 75% pilot-to-production conversion across enterprise engagements by pairing executive strategy with production system delivery.",
            "Author executive-facing AI strategy, agentic systems, MCP, RAG, and frontier-model field analysis for ensolabs.ai and client leadership teams."
        ]
    },
    {
        "company": "Heller Agency",
        "title": "AI Solutions, Center of Excellence & Managed Services Consultant",
        "dates": "2022 - Present | Remote",
        "bullets": [
            "Designed and operate an AI Center of Excellence for a full-service pharma agency across five brand teams under FDA, MLR, and PRC review constraints.",
            "Built five brand knowledge bases, eight active AI automations, deployment playbooks, and governance workflows aligned to NIST AI RMF.",
            "Reduced campaign launch timelines by 83% from three months to two weeks and delivered 35% workflow time savings."
        ]
    },
    {
        "company": "NYC Agency Leadership",
        "title": "VP / Director / Senior Director, Strategy and Experience",
        "dates": "2011 - 2022 | New York, NY",
        "bullets": [
            "Led cross-functional teams of 8-15 across Omnicom, Publicis, WPP, McCann, RAPP, Rokkan, VML/Y&R, and DDB.",
            "Managed $150M+ portfolios for Fortune 500 clients including American Express, Google, Pfizer, Eli Lilly, AT&T, Citi, JPMorgan Chase, and Microsoft.",
            "Delivered enterprise strategy, customer experience, digital transformation, analytics, and growth programs recognized by AdAge and major agency leadership teams."
        ]
    }
]

CORE_BUILDS = {
    "gore": "Architected Gore M2 Intelligence Hub for a Fortune 500 manufacturer: 8-stage LangGraph pipeline, Claude, MCP, Python, AES-256-GCM encryption, 731 documents processed, 16 commercial signals surfaced, April 2026 go/no-go milestone delivered.",
    "google": "Attended Google I/O 2026 as a Google partner; built and analyzed with 150+ developers around Google DeepMind's Managed Agents API, Gemini Omni, Antigravity 2.0, and Gemini 3.5 Flash; translated frontier-model signal into Fortune 500 AI strategy guidance.",
    "trading": "Built and operate Enso Trading Terminal: autonomous signal intelligence and options-trading platform with multi-agent research, Alpaca, Public.com, Hyperliquid, backtesting, options flow, and human-in-the-loop risk controls.",
    "notion": "Shipped open-source career intelligence agent at Notion Developer Platform launch: five Notion Worker tools, TypeScript, multi-model routing, Claude Sonnet/Haiku, MCP integration, and source-tier ranking.",
    "heller": "Built Heller AI Center of Excellence: five brand knowledge bases, eight automations, NIST AI RMF governance, FDA/MLR/PRC-compliant workflows, and 83% faster campaign launches.",
    "site": "Shipped ensolabs.ai as a production Next.js/React site with Vercel deployment, live MCP endpoint, GA4 event tracking, 71 JSON-LD schemas, dynamic Open Graph images, and AI-search optimization."
}

RESUMES = [
    {
        "filename": "Sav_Banerjee_Resume_ForwardDeployed_ATS_2026.docx",
        "headline": "Forward-Deployed AI Architect | Agentic Systems | Enterprise AI Strategy",
        "summary": [
            "Forward-deployed AI architect and enterprise transformation strategist with 15+ years helping Fortune 500 organizations turn AI ambition into production systems.",
            "Operates at the intersection of frontier AI research, executive strategy, and hands-on deployment: agentic workflows, MCP tools, LangGraph, RAG/evaluation, Claude, Gemini, and production operating models.",
            "Known for translating ambiguous executive goals into solution architecture, deployment roadmaps, governance, and measurable value realization."
        ],
        "skills": [
            "Forward-deployed AI strategy", "Agentic AI architecture", "MCP tool integration", "Multi-agent orchestration", "LangGraph", "Claude", "Gemini", "Google Managed Agents API", "RAG and evaluation", "Python", "TypeScript", "React", "Next.js", "APIs", "Cloud deployment", "CI/CD release patterns", "Observability", "C-suite stakeholder alignment", "Pilot-to-production"
        ],
        "proof": [CORE_BUILDS["google"], CORE_BUILDS["gore"], CORE_BUILDS["trading"], CORE_BUILDS["site"], CORE_BUILDS["notion"], CORE_BUILDS["heller"]]
    },
    {
        "filename": "Sav_Banerjee_Resume_AI_Transformation_ATS_2026.docx",
        "headline": "AI Transformation Leader | Enterprise Advisory | Agentic Systems",
        "summary": [
            "Enterprise AI transformation leader with 15+ years guiding Fortune 500 organizations from AI ambition to production.",
            "Combines C-suite advisory, operating-model design, change management, governance, and hands-on agentic system deployment.",
            "Track record includes 75% pilot-to-production conversion, 3-month average time-to-first-value, and $150M+ enterprise portfolios across healthcare, finance, technology, and manufacturing."
        ],
        "skills": [
            "AI transformation strategy", "Enterprise AI roadmap", "Operating model design", "AI readiness assessment", "C-suite workshops", "Responsible AI governance", "Value realization", "Change management", "AI Center of Excellence", "Agentic systems", "RAG", "MCP", "Claude", "Gemini", "Evaluation", "Observability", "Executive enablement"
        ],
        "proof": [CORE_BUILDS["gore"], CORE_BUILDS["heller"], CORE_BUILDS["google"], CORE_BUILDS["site"], CORE_BUILDS["trading"]]
    },
    {
        "filename": "Sav_Banerjee_Resume_MD_ManagedServices_ATS_2026.docx",
        "headline": "AI Managed Services Leader | AI CoE | Implementation Partner",
        "summary": [
            "AI transformation and managed-services leader who designs, deploys, and operates enterprise AI programs after the initial roadmap is approved.",
            "Founder of Enso Labs, a principal-led AI studio where advisory, implementation, governance, and ongoing agent operations are delivered as one integrated practice.",
            "Best fit for AI CoE, implementation partner, deployment lead, managed services, embedded AI operator, and fractional AI leadership roles."
        ],
        "skills": [
            "AI managed services", "AI Center of Excellence", "Implementation partner", "Ongoing agent operations", "Governance playbooks", "Deployment playbooks", "Use-case intake", "Benefits tracking", "Claude managed services", "MCP", "N8N", "RAG knowledge systems", "Evaluation", "Observability", "C-suite executive cohorts", "Pharma compliance", "Value realization"
        ],
        "proof": [CORE_BUILDS["heller"], CORE_BUILDS["gore"], CORE_BUILDS["site"], CORE_BUILDS["notion"], CORE_BUILDS["google"]]
    },
    {
        "filename": "Sav_Banerjee_Resume_PMM_GTM_ATS_2026.docx",
        "headline": "AI Product Strategy | Enterprise GTM | Platform Partnerships",
        "summary": [
            "Product strategy and enterprise go-to-market leader with 15+ years translating complex AI, data, and platform capabilities into market narratives, buyer enablement, and growth programs.",
            "Founder of Enso Labs and Perplexity AI Business Fellowship winner with production work across AI agents, deep research, competitive intelligence, financial AI, and executive demos.",
            "Best fit for AI platform strategy, GTM, product marketing, partnerships, strategic growth, and deployment strategy roles."
        ],
        "skills": [
            "AI product strategy", "Enterprise GTM", "Product marketing", "Positioning", "Buyer enablement", "Competitive intelligence", "Sales enablement", "Strategic partnerships", "AI platform narrative", "Perplexity", "Claude", "Gemini", "MCP", "C-suite executive demos", "Market intelligence"
        ],
        "proof": [CORE_BUILDS["google"], CORE_BUILDS["trading"], CORE_BUILDS["gore"], CORE_BUILDS["site"], CORE_BUILDS["heller"]]
    }
]


def set_cell_text(paragraph, text, bold=False, size=10):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Arial"
    return run


def set_document_defaults(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.05

    for style_name, size in [("Heading 1", 11), ("Heading 2", 10.25)]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 78, 121)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(2)

    bullet = styles["List Bullet"]
    bullet.font.name = "Arial"
    bullet.font.size = Pt(9.3)
    bullet.paragraph_format.left_indent = Inches(0.18)
    bullet.paragraph_format.first_line_indent = Inches(-0.18)
    bullet.paragraph_format.space_after = Pt(1.4)
    bullet.paragraph_format.line_spacing = 1.03


def add_rule(paragraph):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "B7C9D9")
    p_bdr.append(bottom)


def add_header(doc, resume):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    name = p.add_run("SAV BANERJEE")
    name.bold = True
    name.font.name = "Arial"
    name.font.size = Pt(15)
    name.font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    line = p.add_run(resume["headline"])
    line.bold = True
    line.font.name = "Arial"
    line.font.size = Pt(10.2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    set_cell_text(p, CONTACT, size=8.6)
    add_rule(p)


def add_heading(doc, text):
    p = doc.add_paragraph(text.upper(), style="Heading 1")
    p.paragraph_format.keep_with_next = True
    return p


def add_bullets(doc, bullets):
    for item in bullets:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_inline_list(doc, items):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(" | ".join(items))
    run.font.name = "Arial"
    run.font.size = Pt(9.2)


def build_resume(resume):
    doc = Document()
    set_document_defaults(doc)
    add_header(doc, resume)

    add_heading(doc, "Executive Summary")
    for summary in resume["summary"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(summary)
        run.font.name = "Arial"
        run.font.size = Pt(9.5)

    add_heading(doc, "Core Keywords and Capabilities")
    add_inline_list(doc, resume["skills"])

    add_heading(doc, "Selected Production Proof")
    add_bullets(doc, resume["proof"])

    add_heading(doc, "Experience")
    for job in BASE_EXPERIENCE:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        company = p.add_run(job["company"])
        company.bold = True
        company.font.name = "Arial"
        company.font.size = Pt(9.7)
        title = p.add_run(f" | {job['title']}")
        title.bold = True
        title.font.name = "Arial"
        title.font.size = Pt(9.7)
        dates = p.add_run(f" | {job['dates']}")
        dates.font.name = "Arial"
        dates.font.size = Pt(9.0)
        add_bullets(doc, job["bullets"])

    add_heading(doc, "Certifications, Education and Recognition")
    add_bullets(doc, [
        "Claude Certified Architect - Anthropic; Claude Code in Action; Notion AI Agent SDK.",
        "Google AI/Gemini professional certification; Google Cloud devstar program; Google I/O 2026 partner ecosystem access.",
        "Perplexity AI Business Fellowship Winner; 20+ production projects across AI research, agents, and enterprise workflows.",
        "B.A., Advertising - University of Oregon School of Journalism and Communication."
    ])

    add_heading(doc, "Target Role Alignment")
    add_bullets(doc, [
        "Forward-Deployed AI Strategist / Solution Principal; VP or Director AI Transformation; Lead AI Consultant / Principal AI Advisor.",
        "Agentic AI Platform / AI Solutions Architect; AI CoE / Managed Services / Implementation Partner; AI Product Strategy / GTM / Partnerships."
    ])

    out = OUT_DIR / resume["filename"]
    doc.save(out)
    return out


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    written = [build_resume(resume) for resume in RESUMES]
    for path in written:
        print(path)


if __name__ == "__main__":
    main()
